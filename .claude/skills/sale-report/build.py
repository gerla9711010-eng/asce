# -*- coding: utf-8 -*-
"""銷售報告書 產生器 — 套「上銘蒔尚」模板，資料全部從 JSON 讀。

用法:  python build.py <data.json>

模板 13 頁的角色固定，本腳本只換內容、不新增/刪除/重排投影片
（標題旗標與裝飾都是 slide 層級 shape，add_slide() 會得到全白頁）。
"""
import json
import os
import shutil
import sys

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_TICK_MARK
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from helpers import *                                            # noqa: F401,F403

DESKTOP = r'C:\Users\user\OneDrive\桌面'
TEMPLATE = os.path.join(DESKTOP, 'DM', '銷報書', '上銘蒔尚', '@99-銷售報告書.pptx')


def build(spec_path):
    spec = json.load(open(spec_path, encoding='utf-8'))
    src_dir = spec.get('src_dir') or os.path.dirname(os.path.abspath(spec_path))
    name = spec['out_name']
    out_dir = spec.get('out_dir') or os.path.join(DESKTOP, 'DM', '銷報書', name)
    out = os.path.join(out_dir, '%s_銷售報告書.pptx' % name)

    os.makedirs(out_dir, exist_ok=True)
    shutil.copyfile(spec.get('template') or TEMPLATE, out)
    prs = Presentation(out)
    S = prs.slides

    # ---------------------------------------------------------- 1  封面
    c = spec['cover']
    set_runs(find_by_text(S[0], '上銘建設'), [c['area'], '—', c['addr']])
    set_size(find_by_text(S[0], c['addr']), c.get('size', 20))

    # ---------------------------------------------------------- 2  本案現況
    s, d = S[1], spec['subject']
    set_lines(find(s, '文本框 4'), [d.get('badge', '本案現況')])
    drop(find(s, '圖片 7'))
    for nm in ('橢圓 11', '橢圓 12'):
        drop(find(s, nm))

    add_text(s, 0.35, 1.02, 12.6, 0.75, d['title'], size=16, bold=True,
             color=NAVY, align=PP_ALIGN.CENTER, spacing=1.15)
    rows = [['項　目', '內　　容']] + [list(r) for r in d['rows']]
    add_table(s, 0.45, 1.92, 6.35, 4.85, rows, col_w=[1.55, 4.80],
              size=11.5, header_size=12, row_h=4.85 / len(rows),
              align=[PP_ALIGN.CENTER, PP_ALIGN.LEFT])

    notes = find(s, '文字方塊 8')
    move(notes, 7.05, 1.92, 5.85, 3.55)
    set_lines(notes, d['notes'])
    set_size(notes, 14)
    for p in notes.text_frame.paragraphs:          # template colour is too pale
        p.line_spacing = 1.45
        for r in p.runs:
            r.font.color.rgb = NAVY_DEEP
            r.font.bold = True

    # ---------------------------------------------------------- 3  同巷實價
    s, d = S[2], spec['deals']
    set_lines(find(s, '文本框 4'), [d.get('badge', '同巷實價')])
    drop(find(s, '圖片 4'))
    add_text(s, 0.35, 1.02, 12.6, 0.75, d['title'], size=16, bold=True,
             color=NAVY, align=PP_ALIGN.CENTER, spacing=1.15)

    rows = [['交易月份', '門　牌', '樓層', '總價(萬)', '單價(萬/坪)',
             '建物坪數', '持分地坪', '格　局']]
    for r in d['rows']:
        m = '!' if r.get('hl') else ''            # '!' -> red+bold in add_table
        rows.append([m + r['ym'], m + r['addr'], m + r['floor'],
                     m + '{:,}'.format(r['total']), m + '@{:.1f}'.format(r['unit']),
                     m + '{:.1f}坪'.format(r['ping']), m + '{:.1f}坪'.format(r['land']),
                     m + r['layout']])
    add_table(s, 0.85, 1.95, 11.65, 4.15, rows,
              col_w=[1.35, 1.55, 0.90, 1.35, 1.60, 1.50, 1.50, 1.90],
              size=12.5, header_size=12.5, row_h=4.15 / len(rows))
    add_text(s, 0.85, 6.25, 11.65, 1.0, d['footnotes'], size=11.5,
             color=RGBColor(0x55, 0x5F, 0x6E), spacing=1.2)

    # ---------------------------------------------------------- 4  單價走勢
    s, t = S[3], spec['trend']
    set_lines(find(s, '文本框 4'), [t.get('badge', '單價走勢')])
    drop(find(s, '圖片 1'))
    add_text(s, 0.35, 1.02, 12.6, 0.5, [t['title']], size=16, bold=True,
             color=NAVY, align=PP_ALIGN.CENTER)

    chrono = sorted(spec['deals']['rows'], key=lambda r: r['ym'])
    cd = CategoryChartData()
    cd.categories = ['{}\n{}{}'.format(r['ym'], r['addr'], r['floor']) for r in chrono]
    cd.add_series('單價(萬/坪)', tuple(r['unit'] for r in chrono))
    ch = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.45),
                            Inches(1.65), Inches(8.15), Inches(5.3), cd).chart
    set_font_name(ch.font, BODY_FONT)      # PowerPoint ignores the dLbls-level txPr
    ch.font.size = Pt(11)
    ch.has_legend = False
    ch.has_title = False
    plot = ch.plots[0]
    plot.gap_width = 55
    ser = plot.series[0]
    ser.format.fill.solid()
    ser.format.fill.fore_color.rgb = STEEL
    for i, r in enumerate(chrono):         # highlight the same-floor comps
        ser.points[i].format.fill.solid()
        ser.points[i].format.fill.fore_color.rgb = NAVY if r.get('hl') else STEEL
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format = '0.0'
    dl.number_format_is_linked = False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size = Pt(11)
    dl.font.bold = True
    set_font_name(dl.font, BODY_FONT)
    dl.font.color.rgb = NAVY_DEEP
    va = ch.value_axis
    va.minimum_scale = 0.0
    va.maximum_scale = t.get('y_max', max(r['unit'] for r in chrono) * 1.15)
    va.has_major_gridlines = True
    va.tick_labels.font.size = Pt(10)
    set_font_name(va.tick_labels.font, BODY_FONT)
    ca = ch.category_axis
    ca.tick_labels.font.size = Pt(9.5)
    set_font_name(ca.tick_labels.font, BODY_FONT)
    ca.major_tick_mark = XL_TICK_MARK.NONE

    add_text(s, 8.80, 1.70, 4.20, 5.2, t['notes'], size=13, color=INK, spacing=1.3)

    # ---------------------------------------------------------- 5  周遭開價
    s, l = S[4], spec['listings']
    set_lines(find(s, '文本框 4'), [l.get('badge', '周遭開價')])
    drop(find(s, '圖片 1'))
    s.shapes.add_picture(os.path.join(src_dir, l['image']),
                         Inches(0.40), Inches(1.72), Inches(7.30), Inches(5.43))
    add_text(s, 0.35, 1.02, 12.6, 0.55, [l['title']], size=16, bold=True,
             color=NAVY, align=PP_ALIGN.CENTER)
    rows = [['開價區間', '筆數', '明　細（萬）']] + [list(r) for r in l['rows']]
    add_table(s, 7.95, 1.80, 5.00, 2.50, rows, col_w=[1.35, 0.65, 3.00],
              size=12, header_size=12.5, row_h=2.50 / len(rows))
    add_text(s, 7.95, 4.45, 5.00, 2.70, l['notes'], size=12.5, color=INK,
             spacing=1.25)

    # ---------------------------------------------------------- 6  價格建議
    s, p = S[5], spec['price']
    set_lines(find(s, '文本框 4'), [p.get('badge', '價格建議')])
    drop(find(s, '圖片 8'))
    add_text(s, 0.35, 1.02, 12.6, 0.75, p['title'], size=16, bold=True,
             color=NAVY, align=PP_ALIGN.CENTER, spacing=1.15)
    rows = [['定　位', '單價(萬/坪)', '總價區間', '對　照　依　據']] + \
           [list(r) for r in p['rows']]
    add_table(s, 0.60, 1.95, 12.15, 2.05, rows,
              col_w=[1.45, 1.75, 1.90, 7.05], size=13, header_size=13,
              row_h=2.05 / len(rows),
              align=[PP_ALIGN.CENTER, PP_ALIGN.CENTER, PP_ALIGN.CENTER,
                     PP_ALIGN.LEFT])
    add_text(s, 0.60, 4.25, 12.15, 2.85, p['bullets'], size=13, color=INK,
             spacing=1.35)

    band = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.60),
                              Inches(6.20), Inches(12.15), Inches(0.82))
    band.fill.solid()
    band.fill.fore_color.rgb = NAVY
    band.line.fill.background()
    band.shadow.inherit = False
    bp = band.text_frame.paragraphs[0]
    band.text_frame.word_wrap = True
    bp.alignment = PP_ALIGN.CENTER
    br = bp.add_run()
    br.text = p['band']
    set_font_name(br.font, BODY_FONT)
    br.font.size = Pt(15)
    br.font.bold = True
    br.font.color.rgb = WHITE

    # ---------------------------------------------------------- 7  生活圈範圍
    s, lf = S[6], spec['life']
    for box, label in zip(('文本框 11', '文本框 13', '文本框 14'), lf['labels']):
        set_lines(find(s, box), [label])
    for oval, pair in zip(('椭圆 20', '椭圆 17', '椭圆 19'), lf['circle_runs']):
        set_runs(find(s, oval), pair)
    # resolve all three before mutating any — rewriting one can make a later
    # anchor match the wrong box
    boxes = [find_by_text(s, a) for a in
             ('青海路商圈', '緊鄰中都濕地公園', '車程約10分鐘')]
    for box, text in zip(boxes, lf['texts']):
        set_lines(box, [text])

    # ---------------------------------------------------------- 8  交易流程
    if spec.get('flow', {}).get('seller_tax'):
        s = S[7]
        box = [sh for sh in walk(s.shapes)
               if sh.has_text_frame and '完成繳納土增稅' in sh.text_frame.text][0]
        set_lines(box, spec['flow']['seller_tax'])

    # ---------------------------------------------------------- 9  相關稅法
    s, tx = S[8], spec['tax']
    set_runs(find(s, '矩形 29'), tx['box1_label'])
    for nm, top, h, lines in (('文本框 32', 1.48, 1.05, tx['box1']),
                              ('文本框 33', 2.88, 1.75, tx['box2']),
                              ('文本框 34', 5.35, 0.75, tx['box3'])):
        b = find(s, nm)
        move(b, 7.55, top, 5.55, h)        # template boxes overflow at 3+ lines
        set_lines(b, lines)
        set_size(b, 13)
    set_lines(find(s, '文字方塊 7'), tx['circle'])
    if tx.get('note'):
        add_text(s, 7.55, 6.35, 5.40, 0.80, tx['note'], size=11,
                 color=RGBColor(0x55, 0x5F, 0x6E), spacing=1.2)

    prs.save(out)
    return out


if __name__ == '__main__':
    print('saved -> ' + build(sys.argv[1]))
