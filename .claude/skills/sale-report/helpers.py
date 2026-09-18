# -*- coding: utf-8 -*-
"""Shared helpers: edit a pptx while preserving the template's run formatting."""
import copy
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

NAVY      = RGBColor(0x33, 0x3D, 0x50)   # theme tx2 @ lumMod 75% - the badge navy
NAVY_DEEP = RGBColor(0x1F, 0x28, 0x38)
STEEL     = RGBColor(0x7E, 0x97, 0xB8)
BAND      = RGBColor(0xEC, 0xF0, 0xF5)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
INK       = RGBColor(0x2B, 0x2B, 0x2B)
RED       = RGBColor(0xC0, 0x00, 0x00)
BODY_FONT = '微軟正黑體'


# ---------------------------------------------------------------- shape lookup
def walk(shapes):
    for sh in shapes:
        yield sh
        if sh.shape_type == 6:
            for c in walk(sh.shapes):
                yield c


def find(slide, name, nth=0):
    hits = [sh for sh in walk(slide.shapes) if sh.name == name]
    return hits[nth]


def find_by_text(slide, needle):
    for sh in walk(slide.shapes):
        if sh.has_text_frame and needle in sh.text_frame.text:
            return sh
    raise KeyError(needle)


def drop(sh):
    sh._element.getparent().remove(sh._element)


# ----------------------------------------------------------------- font names
def set_font_name(font, name):
    """python-pptx only writes <a:latin>; Chinese needs <a:ea> too or the
    theme's decorative EA typeface takes over."""
    font.name = name
    rPr = font._rPr
    latin = rPr.find(qn('a:latin'))
    if latin is None:
        return
    for tag in ('a:cs', 'a:ea'):          # inserted right after latin -> ea, cs
        old = rPr.find(qn(tag))
        if old is not None:
            rPr.remove(old)
        el = copy.deepcopy(latin)
        el.tag = qn(tag)
        el.set('typeface', name)
        latin.addnext(el)


# ------------------------------------------------- text edits (keep run style)
def _first_run_xml(txBody):
    for p in txBody.findall(qn('a:p')):
        rs = p.findall(qn('a:r'))
        if rs:
            return rs[0]
    return None


def _para_with_runs(txBody):
    for p in txBody.findall(qn('a:p')):
        if p.findall(qn('a:r')):
            return p
    return None


def set_lines(shape, lines):
    """Replace a text frame's content, reusing the template's run properties."""
    tf = shape.text_frame
    txBody = tf._txBody
    proto_p = _para_with_runs(txBody)
    proto_r = _first_run_xml(txBody)

    paras = txBody.findall(qn('a:p'))
    # grow
    while len(paras) < len(lines):
        src = proto_p if proto_p is not None else paras[-1]
        txBody.append(copy.deepcopy(src))
        paras = txBody.findall(qn('a:p'))
    # shrink
    while len(paras) > len(lines):
        txBody.remove(paras[-1])
        paras = txBody.findall(qn('a:p'))

    for p_el, line in zip(paras, lines):
        runs = p_el.findall(qn('a:r'))
        if not runs:
            if proto_r is None:
                continue
            p_el.append(copy.deepcopy(proto_r))
            runs = p_el.findall(qn('a:r'))
        for extra in runs[1:]:
            p_el.remove(extra)
        runs[0].find(qn('a:t')).text = line
        # a break inside the run would survive and duplicate text
        for br in p_el.findall(qn('a:br')):
            p_el.remove(br)


def set_runs(shape, texts):
    """Rewrite the runs of the first run-bearing paragraph, 1:1, keeping style."""
    tf = shape.text_frame
    p_el = _para_with_runs(tf._txBody)
    runs = p_el.findall(qn('a:r'))
    proto = copy.deepcopy(runs[0])
    while len(runs) < len(texts):
        p_el.append(copy.deepcopy(proto))
        runs = p_el.findall(qn('a:r'))
    for extra in runs[len(texts):]:
        p_el.remove(extra)
    runs = p_el.findall(qn('a:r'))
    for r, t in zip(runs, texts):
        r.find(qn('a:t')).text = t


def set_size(shape, pt):
    for p in shape.text_frame.paragraphs:
        for r in p.runs:
            r.font.size = Pt(pt)


def move(shape, l=None, t=None, w=None, h=None):
    if l is not None:
        shape.left = Inches(l)
    if t is not None:
        shape.top = Inches(t)
    if w is not None:
        shape.width = Inches(w)
    if h is not None:
        shape.height = Inches(h)


# ----------------------------------------------------------------- new content
def add_text(slide, l, t, w, h, lines, size=14, bold=False, color=INK,
             align=PP_ALIGN.LEFT, spacing=1.25, font=BODY_FONT):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        r = p.add_run()
        r.text = line
        set_font_name(r.font, font)
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
    return box


def add_table(slide, l, t, w, h, rows, col_w, size=11, header_size=None,
              header_fill=NAVY, row_h=None, align=None):
    """rows[0] is the header. col_w in inches. align: list of PP_ALIGN per col."""
    n_r, n_c = len(rows), len(rows[0])
    gf = slide.shapes.add_table(n_r, n_c, Inches(l), Inches(t),
                                Inches(w), Inches(h))
    tbl = gf.table
    tbl.first_row = False
    tbl.horz_banding = False
    # drop the built-in Office-blue table style
    tblPr = tbl._tbl.find(qn('a:tblPr'))
    if tblPr is not None:
        sid = tblPr.find(qn('a:tableStyleId'))
        if sid is not None:
            tblPr.remove(sid)

    for i, cw in enumerate(col_w):
        tbl.columns[i].width = Inches(cw)
    if row_h:
        for r in tbl.rows:
            r.height = Inches(row_h)

    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = tbl.cell(ri, ci)
            cell.margin_left = Inches(0.05)
            cell.margin_right = Inches(0.05)
            cell.margin_top = Inches(0.02)
            cell.margin_bottom = Inches(0.02)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            if ri == 0:
                cell.fill.fore_color.rgb = header_fill
            else:
                cell.fill.fore_color.rgb = BAND if ri % 2 == 0 else WHITE

            tf = cell.text_frame
            tf.word_wrap = True
            txt = '' if val is None else str(val)
            red = txt.startswith('!')
            if red:
                txt = txt[1:]
            p = tf.paragraphs[0]
            p.alignment = (align[ci] if align else PP_ALIGN.CENTER)
            r = p.add_run()
            r.text = txt
            set_font_name(r.font, BODY_FONT)
            r.font.size = Pt(header_size or size) if ri == 0 else Pt(size)
            r.font.bold = (ri == 0)
            if ri == 0:
                r.font.color.rgb = WHITE
            elif red:
                r.font.color.rgb = RED
                r.font.bold = True
            else:
                r.font.color.rgb = INK
    return gf


def swap_picture(slide, shape, img_path):
    l, t, w, h = shape.left, shape.top, shape.width, shape.height
    drop(shape)
    return slide.shapes.add_picture(img_path, l, t, w, h)
